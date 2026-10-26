#!/usr/bin/env python
"""Generate the synthetic evaluation corpus in sample_submissions/.

Creates 6 PDFs, 6 DOCX files and 3 MP4 clips with known relationships, plus
dataset_manifest.json describing the ground truth. Speech is synthesised
with macOS ``say`` (Indian- and British-English voices); on Linux, install
``espeak-ng`` and the script uses it instead. The generated files are
committed, so re-running this script is only needed to change the corpus.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import dataset_content as C  # noqa: E402

OUT = ROOT / "sample_submissions"


# ------------------------------------------------------------------ PDF helpers
def _pdf(path: Path, team: str, story_fn, two_column: bool = False) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(2 * cm, A4[1] - 1.2 * cm, f"Inter-University Strategy Challenge 2026 | {team}")
        canvas.drawRightString(A4[0] - 2 * cm, 1.0 * cm, f"Confidential - Page {doc.page}")
        canvas.restoreState()

    doc = BaseDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2.2 * cm,
        bottomMargin=2 * cm,
        title=f"{team} submission",
    )
    w, h = A4[0] - 4 * cm, A4[1] - 4.2 * cm
    if two_column:
        gap = 0.6 * cm
        frames = [
            Frame(2 * cm, 2 * cm, (w - gap) / 2, h, id="left"),
            Frame(2 * cm + (w + gap) / 2, 2 * cm, (w - gap) / 2, h, id="right"),
        ]
    else:
        frames = [Frame(2 * cm, 2 * cm, w, h, id="body")]
    doc.addPageTemplates([PageTemplate(id="main", frames=frames, onPage=decorate)])
    doc.build(story_fn())


def _styles():
    from reportlab.lib.styles import getSampleStyleSheet

    s = getSampleStyleSheet()
    s["BodyText"].spaceAfter = 8
    s["BodyText"].leading = 14
    return s


def _prompt_and_refs(styles, with_prompt=True):
    from reportlab.platypus import Paragraph

    head = [Paragraph(C.CASE_PROMPT, styles["Italic"])] if with_prompt else []
    refs = [Paragraph("References", styles["Heading2"])] + [
        Paragraph(r, styles["BodyText"]) for r in C.SHARED_REFERENCES
    ]
    return head, refs


def build_pdfs() -> None:
    from reportlab.lib import colors
    from reportlab.platypus import FrameBreak, PageBreak, Paragraph, Spacer, Table, TableStyle

    st = _styles()
    P = lambda t: Paragraph(t, st["BodyText"])  # noqa: E731
    H = lambda t: Paragraph(t, st["Heading2"])  # noqa: E731

    def sub01():
        head, refs = _prompt_and_refs(st)
        return [
            Paragraph("Team Aarav - VoltRide: Fix Pune First", st["Title"]),
            *head,
            H("1. Situation"),
            P(C.A_SITUATION),
            H("2. Unit economics"),
            P(C.A_ECONOMICS),
            PageBreak(),
            H("3. Recommendation"),
            P(C.A_RECOMMENDATION_X),
            H("4. Risks and mitigation"),
            P(C.A_RISK_Y),
            P(C.A_RISK_Z),
            P(C.A_RISK_3),
            *refs,
        ]

    def sub02():
        head, refs = _prompt_and_refs(st)
        return [
            Paragraph("Team Meridian - Indore First", st["Title"]),
            *head,
            H("Recommendation"),
            P(C.B_INTRO),
            H("Financials"),
            P(C.B_NUMBERS),
            PageBreak(),
            H("Risks"),
            P(C.A_RISK_Y),
            P(C.A_RISK_Z),
            P(C.B_CLOSE),
            *refs,
        ]

    def sub06():
        return [
            Paragraph("Team Ananya - Nagpur Heat Check", st["Title"]),
            H("Background"),
            P(C.WIKI_COPY_1),
            P(C.WIKI_COPY_2),
            H("Our analysis"),
            *[P(t) for t in C.F_OWN],
        ]

    def sub08():
        tbl = Table(C.H_TABLE, colWidths=None)
        tbl.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightblue),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                ]
            )
        )
        return [
            Paragraph("Team Zara - The Coimbatore Case", st["Heading1"]),
            *[P(t) for t in C.H_COL_LEFT],
            FrameBreak(),
            *[P(t) for t in C.H_COL_RIGHT],
            Spacer(1, 8),
            tbl,
            PageBreak(),
            H("Appendix"),
            P(
                "Survey questionnaire available on request. Interviews were "
                "conducted in Tamil and English during the second week of October."
            ),
        ]

    def sub15():
        return [Paragraph("Team Tanvi - Strategic Recommendation", st["Title"]), *[P(t) for t in C.O_AI]]

    _pdf(OUT / "sub_01_aarav.pdf", "Team Aarav", sub01)
    _pdf(OUT / "sub_02_meera.pdf", "Team Meridian", sub02)
    _pdf(OUT / "sub_06_ananya.pdf", "Team Ananya", sub06)
    _pdf(OUT / "sub_08_zara.pdf", "Team Zara", sub08, two_column=True)
    _pdf(OUT / "sub_15_tanvi.pdf", "Team Tanvi", sub15)
    (OUT / "empty_13.pdf").write_bytes(b"")


# ------------------------------------------------------------------ DOCX helpers
def _docx(path: Path, title: str, paragraphs, with_prompt: bool, page_break_after: int = 2) -> None:
    import docx
    from docx.enum.text import WD_BREAK

    d = docx.Document()
    section = d.sections[0]
    section.header.paragraphs[0].text = "Inter-University Strategy Challenge 2026"
    section.footer.paragraphs[0].text = "Confidential"
    d.add_heading(title, level=1)
    if with_prompt:
        d.add_paragraph().add_run(C.CASE_PROMPT).italic = True
    for i, para in enumerate(paragraphs, 1):
        d.add_paragraph(para)
        if i == page_break_after and i < len(paragraphs):
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    if with_prompt:
        d.add_heading("References", level=2)
        for r in C.SHARED_REFERENCES:
            d.add_paragraph(r)
    d.save(str(path))


def build_docx() -> None:
    _docx(OUT / "sub_03_kabir.docx", "Team Kabir - Pilot Before You Scale", C.C_PARAS, True)
    _docx(OUT / "sub_04_diya.docx", "Team Diya - Test One City First", C.D_PARAS, True)
    _docx(OUT / "sub_05_rohan.docx", "Team Rohan - Sell to Employers", C.E_PARAS, True)
    _docx(OUT / "sub_07_ishaan.docx", "Team Ishaan - Strategic Expansion", C.G_AI, False)
    _docx(OUT / "sub_09_vikram.docx", "Team Vikram - Safety Before Scale", C.I_PARAS, False)
    # Degraded: a .docx that is really a truncated zip archive.
    good = OUT / "sub_09_vikram.docx"
    (OUT / "corrupt_14.docx").write_bytes(good.read_bytes()[:700])


# ------------------------------------------------------------------ MP4 helpers
def _ffmpeg() -> str:
    import imageio_ffmpeg

    return imageio_ffmpeg.get_ffmpeg_exe()


def _tts(text: str, wav: Path, voice: str, tempo: float = 1.0) -> None:
    if shutil.which("say"):
        aiff = wav.with_suffix(".aiff")
        subprocess.run(["say", "-v", voice, "-r", "175", "-o", str(aiff), text], check=True)
        # Neural voices ignore `say -r`, so speed is adjusted in ffmpeg instead.
        subprocess.run(
            [_ffmpeg(), "-y", "-loglevel", "error", "-i", str(aiff), "-filter:a", f"atempo={tempo}", str(wav)],
            check=True,
        )
    elif shutil.which("espeak-ng"):
        subprocess.run(["espeak-ng", "-v", "en", "-s", "160", "-w", str(wav), text], check=True)
    else:
        raise SystemExit("Need macOS `say` or `espeak-ng` to synthesise speech")


def _video(out: Path, wav: Path = None, seconds: float = 12.0, color: str = "0x1f3b73") -> None:
    cmd = [_ffmpeg(), "-y", "-loglevel", "error", "-f", "lavfi", "-i", f"color=c={color}:s=640x360:r=15"]
    if wav:
        cmd += ["-i", str(wav), "-shortest", "-c:a", "aac", "-b:a", "96k"]
    else:
        cmd += ["-t", str(seconds)]
    cmd += ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-tune", "stillimage", str(out)]
    subprocess.run(cmd, check=True)


def build_mp4() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        speech = f"{C.PITCH_10_INTRO} {C.A_RECOMMENDATION_X} {C.PITCH_10_OUTRO}"
        _tts(speech, tmp / "p10.wav", "Aman", tempo=1.12)  # Indian-English voice; keeps clip < 30 s
        _video(OUT / "pitch_10_sanjay.mp4", tmp / "p10.wav")
        _tts(C.PITCH_11, tmp / "p11.wav", "Daniel")  # British-English voice
        _video(OUT / "pitch_11_leela.mp4", tmp / "p11.wav", color="0x4a1f73")
    _video(OUT / "silent_12.mp4", None, seconds=15)  # video stream only, no audio track


# ------------------------------------------------------------------ manifest
def manifest() -> dict:
    files = {
        "sub_01_aarav.pdf": {
            "label": "original",
            "ai_generated": False,
            "contains_case_prompt": True,
            "notes": "Source document: paragraphs copied by sub_02 and read aloud in pitch_10.",
        },
        "sub_02_meera.pdf": {"label": "partial_exact_copy", "ai_generated": False, "contains_case_prompt": True},
        "sub_03_kabir.docx": {"label": "original", "ai_generated": False, "contains_case_prompt": True},
        "sub_04_diya.docx": {"label": "paraphrase", "ai_generated": False, "contains_case_prompt": True},
        "sub_05_rohan.docx": {"label": "original_independent", "ai_generated": False, "contains_case_prompt": True},
        "sub_06_ananya.pdf": {"label": "external_copy", "ai_generated": False, "contains_case_prompt": False},
        "sub_07_ishaan.docx": {"label": "ai_generated", "ai_generated": True, "contains_case_prompt": False},
        "sub_08_zara.pdf": {
            "label": "original_heavily_formatted",
            "ai_generated": False,
            "contains_case_prompt": False,
            "notes": "Two-column layout, table, repeated header/footer on every page.",
        },
        "sub_09_vikram.docx": {"label": "original_independent", "ai_generated": False, "contains_case_prompt": False},
        "pitch_10_sanjay.mp4": {
            "label": "cross_modal_copy",
            "ai_generated": False,
            "contains_case_prompt": False,
            "notes": "Indian-English TTS voice reads sub_01 page 2 recommendation verbatim.",
        },
        "pitch_11_leela.mp4": {
            "label": "original",
            "ai_generated": False,
            "contains_case_prompt": False,
            "notes": "British-English TTS voice; original content.",
        },
        "silent_12.mp4": {"label": "degraded_no_audio", "ai_generated": False, "expected_status": "flagged"},
        "empty_13.pdf": {"label": "degraded_zero_byte", "ai_generated": False, "expected_status": "flagged"},
        "corrupt_14.docx": {"label": "degraded_corrupted", "ai_generated": False, "expected_status": "flagged"},
        "sub_15_tanvi.pdf": {"label": "ai_generated", "ai_generated": True, "contains_case_prompt": False},
    }
    for name, f in files.items():
        f.setdefault("expected_status", "ok")
        f["format"] = Path(name).suffix[1:]
    return {
        "description": "Synthetic case-competition corpus (fictional company VoltRide Mobility). All people, "
        "companies and figures are fictional; Wikipedia passages are CC BY-SA 4.0.",
        "passage_note": "source_passages / target_passages name variables in scripts/dataset_content.py; "
        "they give sentence-level ground truth for evaluate.py.",
        "authorship_disclosure": "All texts were drafted with an AI assistant. 'Human' submissions were written to "
        "imitate student prose (specific numbers, first person, uneven sentence length); "
        "'ai_generated' ones use unedited assistant-style prose. AI-detection results on "
        "this corpus therefore measure style separation, not real-world accuracy.",
        "files": files,
        "relationships": [
            {
                "type": "exact_copy",
                "source": "sub_01_aarav.pdf",
                "target": "sub_02_meera.pdf",
                "source_location": {"page": 2},
                "target_location": {"page": 2},
                "source_passages": ["A_RISK_Y", "A_RISK_Z"],
                "target_passages": ["A_RISK_Y", "A_RISK_Z"],
                "description": "Two risk paragraphs copied verbatim.",
            },
            {
                "type": "paraphrase",
                "source": "sub_03_kabir.docx",
                "target": "sub_04_diya.docx",
                "source_passages": ["C_PARAS"],
                "target_passages": ["D_PARAS"],
                "description": "All four argument paragraphs reworded with synonyms and new sentence structure.",
            },
            {
                "type": "cross_modal",
                "source": "sub_01_aarav.pdf",
                "target": "pitch_10_sanjay.mp4",
                "source_location": {"page": 2},
                "source_passages": ["A_RECOMMENDATION_X"],
                "target_passages": ["A_RECOMMENDATION_X"],
                "description": "Recommendation paragraph read aloud in the video pitch.",
            },
            {
                "type": "external_copy",
                "source": "ref:Scooter-sharing system",
                "target": "sub_06_ananya.pdf",
                "source_location": {"url": "https://en.wikipedia.org/wiki/Scooter-sharing_system"},
                "source_passages": [],
                "target_passages": ["WIKI_COPY_1", "WIKI_COPY_2"],
                "description": "Two background paragraphs lifted from Wikipedia.",
            },
        ],
        "boilerplate": {
            "files": [n for n, f in files.items() if f.get("contains_case_prompt")],
            "passage_ids": ["CASE_PROMPT", "SHARED_REFERENCES"],
            "description": "Identical 150-word case prompt plus shared citation list; answers are original "
            "(apart from the relationships above).",
        },
    }


def main() -> None:
    OUT.mkdir(exist_ok=True)
    build_pdfs()
    build_docx()
    build_mp4()
    (OUT / "dataset_manifest.json").write_text(json.dumps(manifest(), indent=2), encoding="utf-8")
    for p in sorted(OUT.iterdir()):
        print(f"{p.name:28s} {p.stat().st_size:>8d} bytes")


if __name__ == "__main__":
    main()
